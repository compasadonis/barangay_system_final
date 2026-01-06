import shutil
import os
import csv
from datetime import datetime, timedelta, timezone
from flask import Flask, render_template, request, redirect, url_for, flash, session, make_response, send_file
from werkzeug.security import generate_password_hash
from sqlalchemy import or_, create_engine
import io
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from models import db, BarangayID, Clearance, Indigency, GoodMoral, FirstJobSeeker, User, ActivityLog
from auth import bp as auth_bp
from config import Config


BASE_DIR = os.path.abspath(os.path.dirname(__file__))
DB_DIR = os.path.join(BASE_DIR, "database")
DB_PATH = os.path.join(DB_DIR, "brgy.db")
BACKUP_DIR = os.path.join(BASE_DIR, "backups")
TEMPLATE_DB = os.path.join(BASE_DIR, "empty_template.db")

os.makedirs(DB_DIR, exist_ok=True)
os.makedirs(BACKUP_DIR, exist_ok=True)

# ---------------------------
# Utility helpers
# ---------------------------
def column_is_date(col):
    return "DATE" in str(col.type).upper() or col.type.__class__.__name__.lower().startswith("date")

def column_is_numeric(col):
    return col.type.__class__.__name__.lower() in ("integer", "bigint", "smallint", "numeric", "float", "decimal")

def model_columns(Model):
    return [c for c in Model.__table__.columns]

def model_headers(Model):
    return [c.name for c in model_columns(Model)]

def make_fields_from_model(Model):
    fields = []
    for c in model_columns(Model):
        # skip internal id column
        if c.name == "id":
            continue

        if column_is_date(c):
            itype = "date"
            default = datetime.utcnow().date().isoformat()
        elif column_is_numeric(c):
            itype = "number"
            default = ""
        else:
            itype = "text"
            default = ""

        options = None
        if c.name == "gender":
            itype = "select"
            options = ["Male", "Female", "Other"]
        elif c.name == "civil_status":
            itype = "select"
            options = ["Single", "Married", "Widowed", "Divorced", "Separated"]
        elif c.name == "years_of_residency" or c.name == "length_of_residency":
            itype = "select"
            options = [str(i) for i in range(1, 51)]

        fd = {
            "name": c.name,
            "type": itype,
            "placeholder": c.name.replace("_", " ").title(),
            "required": not c.nullable and not c.default,
            "col": "col-md-4",
            "options": options,
            "default": default
        }
        fields.append(fd)
    return fields

def row_to_dict(row, Model):
    out = {}
    for c in model_columns(Model):
        val = getattr(row, c.name)
        if isinstance(val, datetime):
            out[c.name] = val.date().isoformat()
        elif hasattr(val, "isoformat") and column_is_date(c) and val is not None:
            try:
                out[c.name] = val.isoformat()
            except:
                out[c.name] = str(val)
        else:
            out[c.name] = val
    return out


def best_display(obj):
    """
    Return a readable display for log records.
    Search priority: name → full_name → username → first_name → resident_name → title → id
    """
    for attr in ("name", "full_name", "username", "first_name", "resident_name", "title"):
        if hasattr(obj, attr) and getattr(obj, attr):
            return f"{getattr(obj, 'id', None)} - {getattr(obj, attr)}"

    if hasattr(obj, "id"):
        return str(obj.id)

    return str(obj)

def log_activity(user, action, table_name, record_id=None):
    """
    Normalize table_name to lowercase, and store record_id as string.
    record_id can be an integer ID or a pre-built display string like "12 - Juan Dela Cruz".
    """
    who = user or session.get("username") or "system"
    tbl = (table_name or "").lower()
    # store record_id as string (so we can save "12 - Name" snapshots)
    rec = record_id if record_id is None else str(record_id)
    log = ActivityLog(user=who, action=action, table_name=tbl, record_id=rec)
    db.session.add(log)
    db.session.commit()

# ---------------------------
# Template DB generator
# ---------------------------
def generate_empty_template():
    """
    Create empty_template.db (SQLite) with tables from models.db.metadata
    if it doesn't already exist.
    """
    if os.path.exists(TEMPLATE_DB):
        return

    print("No empty_template.db found — creating one automatically...")


    eng = create_engine(f"sqlite:///{TEMPLATE_DB}")
    db.metadata.create_all(bind=eng)

    print("empty_template.db created successfully at:", TEMPLATE_DB)

# ---------------------------
# Initialize DB defaults 
# ---------------------------
def init_db(app):
    with app.app_context():
        db.create_all()

        # create default users if not exist
        if not User.query.filter_by(username="captain").first():
            db.session.add(User(
                username="captain",
                password=generate_password_hash("captain123"),
                role="admin"
            ))

        if not User.query.filter_by(username="secretary").first():
            db.session.add(User(
                username="secretary",
                password=generate_password_hash("secretary123"),
                role="staff"
            ))

        db.session.commit()

# ---------------------------
# Create app
# ---------------------------
def create_app():
    app = Flask(__name__, static_folder="static", template_folder="templates")

    app.config.from_object(Config)

    if not app.config.get("SQLALCHEMY_DATABASE_URI"):
        app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///" + DB_PATH
    app.config.setdefault("SQLALCHEMY_TRACK_MODIFICATIONS", False)

    db.init_app(app)

    generate_empty_template()

    init_db(app)

    try:
        app.register_blueprint(auth_bp)
    except Exception:
        pass

    app.jinja_env.globals["getattr"] = getattr

    # ---------------------------
    # Forgot Password 
    # ---------------------------
    @app.route("/forgot_password", methods=["GET", "POST"])
    def forgot_password():
        if request.method == "POST":
            username = request.form.get("username")

            user = User.query.filter_by(username=username).first()
            if not user:
                flash("User not found!", "danger")
                return redirect(url_for("forgot_password"))

            # Reset password
            new_password = f"{username}123"
            user.password = generate_password_hash(new_password)
            db.session.commit()

            log_activity(
                session.get("username"),
                "RESET PASSWORD",
                "users",
                user.id
            )

            flash(f"Password reset to: {new_password}", "success")
            return redirect(url_for("auth.login"))

        # GET request — fetch all users
        users = User.query.all()  # <-- Make sure this line is here
        return render_template("forgot_password.html", users=users)


    # ---------------------------
    # Routes (dashboard)
    # ---------------------------
    @app.route("/")
    def index():
        role = session.get("role")
        if role == "admin":
            return redirect(url_for("admin_dashboard"))
        if role == "staff":
            return redirect(url_for("staff_dashboard"))
        return redirect(url_for("auth.login"))

    @app.route("/admin")
    def admin_dashboard():
        if session.get("role") != "admin":
            return redirect(url_for("auth.login"))

        counts = {
            "barangay_id": BarangayID.query.count(),
            "clearance": Clearance.query.count(),
            "indigency": Indigency.query.count(),
            "goodmoral": GoodMoral.query.count(),
            "firstjob": FirstJobSeeker.query.count()
        }

        return render_template("dashboard_admin.html", counts=counts)

    @app.route("/staff")
    def staff_dashboard():
        if session.get("role") != "staff":
            return redirect(url_for("auth.login"))
        return render_template("dashboard_secretary.html")

    # -------------------------------------------------------
    # Generic register_routes for modules
    # -------------------------------------------------------
    def register_routes(prefix, Model, template, title):
        endpoint_name = f"view_{prefix}"

        def view_func():
            # ---------------------------
            # POST: Create Record
            # ---------------------------
            if request.method == "POST":
                required_fields = [
                    c.name for c in Model.__table__.columns
                    if not c.nullable and c.name != "id"
                ]

                for rf in required_fields:
                    if not request.form.get(rf):
                        flash(f"'{rf.replace('_',' ').title()}' is required!", "danger")
                        return redirect(request.url)

                name = request.form.get("name")
                if name and Model.query.filter_by(name=name).first():
                    flash("This person already exists in the records!", "danger")
                    return redirect(request.url)

                obj = Model()
                data = request.form.to_dict()

                for col in Model.__table__.columns:
                    colname = col.name
                    if colname == "id":
                        continue

                    val = data.get(colname)


                    if val and column_is_date(col):
                        try:
                            val = datetime.strptime(val, "%Y-%m-%d").date()
                        except:
                            try:
                                val = datetime.fromisoformat(val)
                            except:
                                val = None

                    if val and column_is_numeric(col):
                        try:
                            if col.type.__class__.__name__.lower() in ("integer", "bigint", "smallint"):
                                val = int(val)
                            else:
                                val = float(val)
                        except:
                            pass

                    setattr(obj, colname, val)

                # Default status
                if hasattr(obj, "status") and not obj.status:
                    obj.status = "Valid"

                db.session.add(obj)
                db.session.commit()

                log_activity(
                    session.get("username"),
                    "CREATE",
                    Model.__tablename__ or prefix,
                    best_display(obj)
                )

                flash(f"{title} created successfully!", "success")
                return redirect(request.url)

            # ---------------------------
            # GET: Fetch Records
            # ---------------------------
            q = request.args.get("q", "").strip()
            month = request.args.get("month", "")
            year = request.args.get("year", "")

            qry = Model.query

            # Search
            if q:
                like = f"%{q}%"
                text_cols = [
                    getattr(Model, c.name)
                    for c in Model.__table__.columns
                    if hasattr(c.type, "length")
                    or c.type.__class__.__name__.lower() in ("text", "varchar", "string")
                ]

                if text_cols:
                    qry = qry.filter(or_(*[col.ilike(like) for col in text_cols]))

            if month and year and "date_issued" in Model.__table__.columns.keys():
                qry = qry.filter(
                    db.func.strftime("%m", Model.date_issued) == f"{int(month):02d}",
                    db.func.strftime("%Y", Model.date_issued) == year
                )

            records_orm = qry.order_by(Model.id.desc()).all()

            today = datetime.utcnow().date()
            for r in records_orm:
                if hasattr(r, "date_issued") and r.date_issued:
                    months = 12 if "business" in (getattr(r, "purpose", "") or "").lower() else 6
                    expiry = r.date_issued + timedelta(days=30 * months)

                    if hasattr(r, "status"):
                        r.status = "Expired" if today > expiry else "Valid"

            records = [row_to_dict(r, Model) for r in records_orm]
            headers = model_headers(Model)
            fields = make_fields_from_model(Model)

            # ---------------------------
            # PAGINATION
            # ---------------------------
            page = request.args.get("page", 1, type=int)
            per_page = 10
            total = len(records)

            start = (page - 1) * per_page
            end = start + per_page
            page_records = records[start:end]

            class PageObj:
                def __init__(self, page, per_page, total):
                    self.page = page
                    self.per_page = per_page
                    self.total = total
                    self.pages = (total + per_page - 1) // per_page

            page_obj = PageObj(page, per_page, total)

            # ---------------------------
            # Render template
            # ---------------------------
            return render_template(
                template,
                title=title,
                records=page_records,
                page_obj=page_obj,
                q=q,
                month=month,
                year=year,
                route_name=prefix,
                headers=headers,
                fields=fields
            )

        app.add_url_rule(
            f"/{prefix}",
            endpoint=endpoint_name,
            view_func=view_func,
            methods=["GET", "POST"]
        )

    # register module routes
    register_routes("barangay_id", BarangayID, "barangay_id.html", "Barangay ID")
    register_routes("clearance", Clearance, "clearance.html", "Clearance")
    register_routes("indigency", Indigency, "indigency.html", "Indigency")
    register_routes("goodmoral", GoodMoral, "good_moral.html", "Good Moral")
    register_routes("firstjob", FirstJobSeeker, "job_seeker.html", "First Job Seeker")

    # ---------------------------
    # Edit / Delete / Print
    # ---------------------------
    @app.route("/<rtype>/edit/<int:id>", methods=["GET","POST"])
    def edit_record(rtype,id):
        mapping = {"barangay_id":BarangayID,"clearance":Clearance,"indigency":Indigency,"goodmoral":GoodMoral,"firstjob":FirstJobSeeker}
        Model = mapping.get(rtype)
        if not Model:
            flash("Invalid type","danger")
            return redirect(url_for("index"))
        record = Model.query.get_or_404(id)
        if request.method=="POST":
            for col in Model.__table__.columns:
                if col.name=="id": continue
                val = request.form.get(col.name)
                if val and column_is_date(col):
                    try: val=datetime.strptime(val,"%Y-%m-%d").date()
                    except: pass
                if val and column_is_numeric(col):
                    try:
                        if col.type.__class__.__name__.lower() in ("integer","bigint","smallint"): val=int(val)
                        else: val=float(val)
                    except: pass
                setattr(record,col.name,val)
            db.session.commit()

            log_activity(
                session.get("username"),
                "UPDATE",
                Model.__tablename__ or rtype,
                best_display(record)
            )

            flash("Updated successfully!","success")
            return redirect(url_for(f"view_{rtype}"))
        record_dict=row_to_dict(record,Model)
        columns=model_columns(Model)
        headers=model_headers(Model)
        fields=make_fields_from_model(Model)
        return render_template("edit_generic.html",record=record_dict,columns=columns,rtype=rtype,headers=headers,fields=fields)

    @app.route("/<rtype>/delete/<int:id>", methods=["POST"])
    def delete_record(rtype, id):
        mapping = {
            "barangay_id": BarangayID,
            "clearance": Clearance,
            "indigency": Indigency,
            "goodmoral": GoodMoral,
            "firstjob": FirstJobSeeker
        }
        Model = mapping.get(rtype)
        if not Model:
            flash("Invalid type", "danger")
            return redirect(url_for("index"))

        record = Model.query.get_or_404(id)


        record_display = best_display(record)

        # delete
        db.session.delete(record)
        db.session.commit()

        # log with the snapshot display (string) so it persists in the audit trail
        log_activity(session.get("username"), "DELETE", Model.__tablename__ or rtype, record_display)

        flash("Record deleted!", "success")
        return redirect(url_for(f"view_{rtype}"))

    @app.route("/print")
    def print_view():
        rtype = request.args.get("rtype")
        month = request.args.get("month")
        year = request.args.get("year")
        mapping = {"barangay_id":(BarangayID,"Barangay ID"),"clearance":(Clearance,"Clearance"),"indigency":(Indigency,"Indigency"),"goodmoral":(GoodMoral,"Good Moral"),"firstjob":(FirstJobSeeker,"First Job Seeker")}
        if rtype not in mapping:
            flash("Invalid print type","danger")
            return redirect(url_for("index"))
        Model,title = mapping[rtype]
        qry = Model.query
        if month and year and "date_issued" in Model.__table__.columns.keys():
            qry = qry.filter(db.func.strftime("%m",Model.date_issued)==f"{int(month):02d}",db.func.strftime("%Y",Model.date_issued)==year)
        records_orm = qry.order_by(Model.id.asc()).all()
        today=datetime.utcnow().date()
        for r in records_orm:
            if hasattr(r,"date_issued") and r.date_issued:
                months=12 if "business" in (getattr(r,"purpose","") or "").lower() else 6
                expiry=r.date_issued + timedelta(days=30*months)
                if hasattr(r,"status"): r.status="Expired" if today>expiry else "Valid"
        records=[row_to_dict(r,Model) for r in records_orm]

        if request.args.get("export")=="csv":
            import io
            si=io.StringIO()
            writer=csv.writer(si)
            cols=[c.name for c in Model.__table__.columns]
            writer.writerow(cols)
            for r in records:
                row_vals=[]
                for c in cols:
                    v=r.get(c)
                    if isinstance(v,(datetime,)):
                        row_vals.append(v.date().isoformat())
                    else:
                        row_vals.append("" if v is None else str(v))
                writer.writerow(row_vals)
            resp=make_response(si.getvalue())
            resp.headers["Content-Disposition"]=f"attachment; filename={rtype}_{month or 'all'}_{year or 'all'}.csv"
            resp.headers["Content-Type"]="text/csv"
            return resp

        return render_template("printable.html",records=records,title=title,month=month,year=year)

    # ---------------------------
    # Activity Log
    # ---------------------------
    @app.route("/activity_log")
    def activity_log_view():
        if session.get("role") not in ("admin", "staff"):
            return redirect(url_for("auth.login"))

        from sqlalchemy import or_

        q = request.args.get("q", "").strip()
        base_query = ActivityLog.query


        if q:
            base_query = base_query.filter(
                or_(
                    ActivityLog.user.ilike(f"%{q}%"),
                    ActivityLog.action.ilike(f"%{q}%"),
                    ActivityLog.table_name.ilike(f"%{q}%"),
                    ActivityLog.record_id.ilike(f"%{q}%")
                )
            )


        page = int(request.args.get("page", 1))
        per_page = 10
        total_logs = base_query.count()

        logs = (
            base_query.order_by(ActivityLog.timestamp.desc())
            .offset((page - 1) * per_page)
            .limit(per_page)
            .all()
        )

        gmt8 = timezone(timedelta(hours=8))

        table_map = {
            "barangay_id": BarangayID,
            "clearance": Clearance,
            "indigency": Indigency,
            "goodmoral": GoodMoral,
            "good_moral": GoodMoral,
            "firstjob": FirstJobSeeker,
            "first_job_seeker": FirstJobSeeker,
            "first_job": FirstJobSeeker,
            "users": User,
            "user": User
        }

        def resolve_model(table_name):
            """Match table_name to model even with s/plurals."""
            if not table_name:
                return None
            t = table_name.lower().strip()
            if t in table_map:
                return table_map[t]
            t2 = t.rstrip("s")
            for key, M in table_map.items():
                if key.rstrip("s") == t2:
                    return M
            return None

        def get_record_display_name(record):
            """Get readable record name."""
            if not record:
                return None
            for attr in ("name", "full_name", "username", "first_name", "resident_name", "title"):
                if hasattr(record, attr) and getattr(record, attr):
                    return getattr(record, attr)
            if hasattr(record, "id"):
                return f"ID {record.id}"
            return None


        for log in logs:

            # Timestamp → GMT+8
            if log.timestamp:
                if log.timestamp.tzinfo is None:
                    log.timestamp = log.timestamp.replace(tzinfo=timezone.utc)
                log.timestamp = log.timestamp.astimezone(gmt8)

            Model = resolve_model(log.table_name)


            if log.record_id and not str(log.record_id).isdigit():
                log.record_name = log.record_id
                continue


            if log.table_name and log.table_name.lower().rstrip("s") == "user":

                if log.record_id:
                    try:
                        uid = int(str(log.record_id).split(" ")[0])
                        user_obj = db.session.get(User, uid)
                    except:
                        user_obj = None

                    if user_obj:
                        log.record_name = f"{user_obj.id} - {user_obj.username}"
                    else:
                        # user deleted – still show ID
                        log.record_name = str(log.record_id)
                else:
                    # fallback: show who performed the action (very rare)
                    log.record_name = log.user

                continue


            if Model and log.record_id:
                try:
                    clean_id = int(str(log.record_id).split(" ")[0])
                    record = db.session.get(Model, clean_id)
                except:
                    record = None

                if record:
                    rn = get_record_display_name(record)
                    log.record_name = rn if rn else str(log.record_id)
                else:
                    # record deleted – still show ID
                    log.record_name = str(log.record_id)
            else:
                # fallback
                log.record_name = str(log.record_id)


        class PageObj:
            def __init__(self, page, per_page, total):
                self.page = page
                self.per_page = per_page
                self.total = total

            @property
            def has_previous(self):
                return self.page > 1

            @property
            def has_next(self):
                return self.page * self.per_page < self.total

            @property
            def previous_page_number(self):
                return self.page - 1

            @property
            def next_page_number(self):
                return self.page + 1

            @property
            def num_pages(self):
                from math import ceil
                return ceil(self.total / self.per_page)

        page_obj = PageObj(page, per_page, total_logs)

        return render_template(
            "activity_log.html",
            logs=logs,
            page_obj=page_obj
        )

    # ---------------------------
    # Export logs to Excel
    # ---------------------------
    @app.route("/export-logs-excel")
    def export_logs_excel():
        logs = ActivityLog.query.order_by(ActivityLog.timestamp.desc()).all()

        wb = Workbook()
        ws = wb.active
        ws.title = "Activity Logs"

        ws.append(["ID", "User", "Action", "Table", "Record", "Timestamp"])

        # same robust mapping used above
        table_map_candidates = {
            "barangay_id": BarangayID,
            "clearance": Clearance,
            "indigency": Indigency,
            "goodmoral": GoodMoral,
            "good_moral": GoodMoral,
            "firstjob": FirstJobSeeker,
            "first_job_seeker": FirstJobSeeker,
            "user": User,
            "users": User
        }

        def resolve_model(table_name):
            if not table_name:
                return None
            t = table_name.lower()
            if t in table_map_candidates:
                return table_map_candidates[t]
            t_stripped = t.rstrip("s")
            for key, M in table_map_candidates.items():
                if key in t or t in key or key.rstrip("s") == t_stripped:
                    return M
            return None

        def get_record_display_name_for_export(record):
            if not record:
                return ""
            for attr in ("name", "full_name", "username", "first_name", "title"):
                if hasattr(record, attr) and getattr(record, attr):
                    return getattr(record, attr)
            if hasattr(record, "id"):
                return str(getattr(record, "id"))
            return ""

        for log in logs:
            ts = ""
            if log.timestamp:
                if log.timestamp.tzinfo is None:
                    log.timestamp = log.timestamp.replace(tzinfo=timezone.utc)
                ts = log.timestamp.astimezone(
                    timezone(timedelta(hours=8))
                ).strftime("%Y-%m-%d %H:%M:%S")

            Model = resolve_model(log.table_name)
            if Model and log.record_id:
                try:
                    record = db.session.get(Model, int(log.record_id))
                except Exception:
                    record = db.session.get(Model, log.record_id)
                record_name = get_record_display_name_for_export(record)
            else:
                record_name = ""

            ws.append([
                log.id,
                log.user,
                log.action,
                log.table_name,
                record_name,
                ts
            ])

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"activity_logs_{datetime.now().strftime('%Y-%m-%d')}.xlsx"

        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    @app.route("/system_settings")
    def system_settings():
        return render_template("system_settings.html")

    # ---------------------------
    # Account Settings 
    # ---------------------------
    @app.route("/account_settings")
    def account_settings():
        if session.get("role") != "admin":
            flash("Admins only!", "danger")
            return redirect(url_for("index"))

        users = User.query.order_by(User.id.asc()).all()
        return render_template("account_settings.html", users=users)

    @app.route("/account/add", methods=["POST"])
    def add_user():
        if session.get("role") != "admin":
            return redirect(url_for("index"))

        username = request.form.get("username")
        password = request.form.get("password")
        role = request.form.get("role")

        if User.query.filter_by(username=username).first():
            flash("Username already exists!", "danger")
            return redirect(url_for("account_settings"))

        new_user = User(
            username=username,
            password=generate_password_hash(password),
            role=role
        )
        db.session.add(new_user)
        db.session.commit()

        log_activity(session.get("username"), "CREATE USER", "users", best_display(new_user))
        flash("User created!", "success")
        return redirect(url_for("account_settings"))

    @app.route("/account/update_password", methods=["POST"])
    def update_password():
        username = request.form.get("username")
        new_password = request.form.get("new_password")
        confirm_password = request.form.get("confirm_password")

        if new_password != confirm_password:
            flash("Passwords do not match!", "danger")
            return redirect(url_for("account_settings"))

        user = User.query.filter_by(username=username).first()
        if not user:
            flash("User not found!", "danger")
            return redirect(url_for("account_settings"))

        user.password = generate_password_hash(new_password)
        db.session.commit()

        log_activity(session.get("username"), "CHANGE PASSWORD", "users", best_display(user))
        flash("Password updated successfully!", "success")
        return redirect(url_for("account_settings"))

    @app.route("/account/reset/<int:user_id>", methods=["POST"])
    def reset_user_password(user_id):
        if session.get("role") != "admin":
            return redirect(url_for("index"))

        user = User.query.get_or_404(user_id)
        user.password = generate_password_hash("password123")
        db.session.commit()

        log_activity(session.get("username"), "RESET PASSWORD", "users", best_display(user))
        flash(f"Password reset to default: password123", "warning")
        return redirect(url_for("account_settings"))

    @app.route("/account/delete/<int:user_id>", methods=["POST"])
    def delete_user(user_id):
        if session.get("role") != "admin":
            return redirect(url_for("index"))

        user = User.query.get_or_404(user_id)


        user_display = f"{user.id} - {user.username}" if getattr(user, "username", None) else f"{user.id}"

        db.session.delete(user)
        db.session.commit()

        # log using table name "users" to match conventions
        log_activity(session.get("username"), "DELETE USER", "users", user_display)
        flash("User deleted!", "danger")
        return redirect(url_for("account_settings"))

    # ---------------------------
    # Backup / Restore / Reset
    # ---------------------------

    def get_db_path():
        uri = app.config["SQLALCHEMY_DATABASE_URI"]
        if uri.startswith("sqlite:///"):
            return os.path.abspath(uri.replace("sqlite:///", ""))
        return None

    @app.route("/backup_recovery")
    def backup_recovery():
        if session.get("role") != "admin":
            flash("Admins only!", "danger")
            return redirect(url_for("index"))
        return render_template("backup_recovery.html")

    @app.route("/backup_database")
    def backup_database():
        try:
            db_path = get_db_path()

            if not os.path.exists(db_path):
                raise FileNotFoundError("Database file not found!")

            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            backup_file = os.path.join(BACKUP_DIR, f"brgy_backup_{timestamp}.db")

            shutil.copy(db_path, backup_file)

            flash("Backup created!", "success")
            return send_file(backup_file, as_attachment=True)

        except Exception as e:
            flash(f"Backup failed: {e}", "danger")
            return redirect(url_for("backup_recovery"))

    @app.route("/restore_database", methods=["POST"])
    def restore_database():
        try:
            file = request.files.get("db_file")
            if not file or not file.filename.endswith(".db"):
                flash("Invalid file!", "danger")
                return redirect(url_for("backup_recovery"))

            temp_path = "temp_restore.db"
            file.save(temp_path)

            import sqlite3

            backup_conn = sqlite3.connect(temp_path)
            backup_cursor = backup_conn.cursor()

            live_db_path = get_db_path()
            live_conn = sqlite3.connect(live_db_path)
            live_cursor = live_conn.cursor()

            protected_tables = ["user"]

            tables = backup_cursor.execute(
                "SELECT name FROM sqlite_master WHERE type='table'"
            ).fetchall()

            for (table_name,) in tables:
                if table_name.lower() in protected_tables:
                    continue

                live_cursor.execute(f"DELETE FROM {table_name}")

                rows = backup_cursor.execute(f"SELECT * FROM {table_name}").fetchall()

                if rows:
                    placeholders = ",".join("?" * len(rows[0]))
                    live_cursor.executemany(
                        f"INSERT INTO {table_name} VALUES ({placeholders})", rows
                    )

            live_conn.commit()
            backup_conn.close()
            live_conn.close()

            os.remove(temp_path)

            flash("Database restored successfully! (Users preserved)", "success")
            return redirect(url_for("backup_recovery"))

        except Exception as e:
            flash(f"Restore failed: {e}", "danger")
            return redirect(url_for("backup_recovery"))

    @app.route("/reset_database", methods=["POST"])
    def reset_database():
        try:
            db.session.rollback()
            protected_tables = ["user"]

            meta = db.metadata

            for table in reversed(meta.sorted_tables):
                if table.name.lower() not in protected_tables:
                    db.session.execute(table.delete())

            db.session.commit()

            flash("System reset successfully! (Users preserved)", "warning")
            return redirect(url_for("backup_recovery"))

        except Exception as e:
            db.session.rollback()
            flash(f"Reset failed: {e}", "danger")
            return redirect(url_for("backup_recovery"))

    return app

# ---------------------------
# Run the app
# ---------------------------
if __name__ == "__main__":
    app = create_app()
    app.run(debug=True)
